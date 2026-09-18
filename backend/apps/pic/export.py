"""Exportación Excel — seguimiento PIC."""
from __future__ import annotations

from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

from apps.entities.models import Entity

from .access import actividades_queryset_for_user
from .calculos import actividad_metrics
from .models import Trimestre

HEADER_FILL = PatternFill(start_color="0E7490", end_color="0E7490", fill_type="solid")
HEADER_FONT = Font(bold=True, color="FFFFFF", size=11)
THIN_BORDER = Border(
    left=Side(style="thin"),
    right=Side(style="thin"),
    top=Side(style="thin"),
    bottom=Side(style="thin"),
)

COLUMNS = [
    "NÚMERO",
    "EJE ESTRATÉGICO",
    "LÍNEA OPERATIVA",
    "ENCARGADO",
    "ACTIVIDAD",
    "UNIDAD",
    "TOTAL PROGRAMADO",
    "EJECUTADO",
    "DISPONIBLE",
    "AVANCE %",
    "VALOR TOTAL",
    "VALOR COBRADO",
    "T1 PROG",
    "T2 PROG",
    "T3 PROG",
    "T4 PROG",
    "RESPONSABLES",
]


def build_seguimiento_excel(
    user,
    entity: Entity,
    *,
    anio: int,
    trimestre: int | None = None,
) -> tuple[BytesIO, str]:
    act_qs = actividades_queryset_for_user(user, entity).filter(plan__anio=anio).order_by("numero")
    if trimestre:
        field_map = {1: "prog_t1", 2: "prog_t2", 3: "prog_t3", 4: "prog_t4"}
        f = field_map.get(trimestre)
        if f:
            act_qs = act_qs.filter(**{f"{f}__gt": 0})

    wb = Workbook()
    ws = wb.active
    ws.title = "Seguimiento PIC"

    for col_idx, title in enumerate(COLUMNS, start=1):
        cell = ws.cell(row=1, column=col_idx, value=title)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = THIN_BORDER

    row_idx = 2
    for act in act_qs.prefetch_related("responsables"):
        m = actividad_metrics(act)
        responsables = ", ".join(u.get_full_name() or u.email for u in act.responsables.all())
        values = [
            act.numero,
            act.eje_estrategico,
            act.linea_operativa,
            act.encargado_texto,
            act.actividad,
            act.unidad_medida,
            act.total_programado,
            m["total_ejecutado"],
            m["disponible"],
            m["avance_pct"],
            float(act.valor_total),
            float(m["valor_cobrado_total"]),
            act.prog_t1,
            act.prog_t2,
            act.prog_t3,
            act.prog_t4,
            responsables,
        ]
        for col_idx, val in enumerate(values, start=1):
            cell = ws.cell(row=row_idx, column=col_idx, value=val)
            cell.border = THIN_BORDER
            cell.alignment = Alignment(wrap_text=True, vertical="top")
        row_idx += 1

    buf = BytesIO()
    wb.save(buf)
    buf.seek(0)
    tri_suffix = f"_T{trimestre}" if trimestre else ""
    filename = f"PIC_{entity.slug or entity.id}_{anio}{tri_suffix}.xlsx"
    return buf, filename
