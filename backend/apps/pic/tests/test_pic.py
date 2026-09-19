"""Pruebas del módulo PIC."""
from __future__ import annotations

from decimal import Decimal
from io import BytesIO

from django.contrib.auth.models import Group
from django.core.files.uploadedfile import SimpleUploadedFile
from django.test import TestCase
from rest_framework.test import APIRequestFactory, force_authenticate

from apps.accounts.models import User, UserEntityMembership
from apps.entities.models import Entity, Secretaria
from apps.pic.calculos import calcular_valor_cobrado, calcular_valor_unitario, trimestre_from_date
from apps.pic.excel_import import import_pic_excel, normalizar_etiqueta, parse_pic_excel, tokenizar_encargado
from apps.pic.models import PicActividad, PicCargo, PicEjecucion, PicPlan
from apps.pic.validators import validate_evidencia_archivo
from apps.pic.views import PicActividadViewSet
from rest_framework.exceptions import ValidationError


def _minimal_xlsx_rows(rows: list[dict]) -> bytes:
    from openpyxl import Workbook

    wb = Workbook()
    ws = wb.active
    ws["A1"] = "EJE"
    ws["C1"] = "Numero"
    ws["D1"] = "ENCARGADO"
    ws["E1"] = "ACTIVIDAD"
    ws["F1"] = "SOPORTES"
    ws["G1"] = "UNIDAD"
    ws["H1"] = "TOTAL"
    ws["I1"] = "I"
    ws["J1"] = "II"
    ws["K1"] = "III"
    ws["L1"] = "IV"
    ws["M1"] = "VALOR"
    for idx, row in enumerate(rows, start=3):
        ws[f"C{idx}"] = row["numero"]
        ws[f"D{idx}"] = row.get("encargado", "ENFERMERA")
        ws[f"E{idx}"] = row.get("actividad", f"Actividad {row['numero']}")
        ws[f"F{idx}"] = row.get("soportes", "PDF soporte")
        ws[f"G{idx}"] = row.get("unidad", "Unidad")
        ws[f"H{idx}"] = row["total"]
        ws[f"I{idx}"] = row.get("t1", 0)
        ws[f"J{idx}"] = row.get("t2", 0)
        ws[f"K{idx}"] = row.get("t3", 0)
        ws[f"L{idx}"] = row.get("t4", 0)
        ws[f"M{idx}"] = row.get("valor", 12000000)
    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()


class PicCalculosTests(TestCase):
    def test_valor_unitario(self):
        self.assertEqual(calcular_valor_unitario(Decimal("12000000"), 4), Decimal("3000000.00"))

    def test_valor_cobrado_suma_exacta(self):
        class Act:
            valor_total = Decimal("12000000")
            total_programado = 4

        act = Act()
        v1 = calcular_valor_cobrado(act, 0, 1)
        v2 = calcular_valor_cobrado(act, 1, 1)
        v3 = calcular_valor_cobrado(act, 2, 1)
        v4 = calcular_valor_cobrado(act, 3, 1)
        self.assertEqual(v1 + v2 + v3 + v4, Decimal("12000000.00"))

    def test_trimestre_from_date(self):
        from datetime import date

        self.assertEqual(trimestre_from_date(date(2026, 2, 10), 2026), 1)
        self.assertEqual(trimestre_from_date(date(2026, 8, 10), 2026), 3)


class PicValidatorTests(TestCase):
    def test_rejects_non_pdf(self):
        with self.assertRaises(ValidationError):
            validate_evidencia_archivo("foto.jpg", 100)

    def test_accepts_pdf(self):
        validate_evidencia_archivo("soporte.pdf", 1024)


class PicExcelImportTests(TestCase):
    def setUp(self):
        self.entity = Entity.objects.create(name="Ent PIC", code="PIC", slug="ent-pic", enable_pic=True)
        self.admin = User.objects.create_user(email="admin-pic@test.com", password="x", entity=self.entity)
        self.admin.groups.add(Group.objects.create(name="admin"))
        UserEntityMembership.objects.create(user=self.admin, entity=self.entity, role="admin", is_active=True)

    def test_rejects_duplicate_numero(self):
        content = _minimal_xlsx_rows(
            [
                {"numero": 1, "total": 4, "t1": 1, "t2": 1, "t3": 1, "t4": 1},
                {"numero": 1, "total": 2, "t1": 1, "t2": 1},
            ]
        )
        rows, errores = parse_pic_excel(content)
        self.assertTrue(any("duplicado" in e.lower() for e in errores))

    def test_import_preserves_ejecuciones_on_reupload(self):
        content = _minimal_xlsx_rows([{"numero": 1, "total": 4, "t1": 1, "t2": 1, "t3": 1, "t4": 1}])
        r1 = import_pic_excel(self.entity, 2026, content, "pic.xlsx", self.admin)
        self.assertTrue(r1["ok"])
        act = PicActividad.objects.get(plan__anio=2026, numero=1)
        PicEjecucion.objects.create(
            entity=self.entity,
            actividad=act,
            fecha_ejecucion="2026-02-15",
            trimestre=1,
            cantidad_ejecutada=1,
            valor_cobrado=Decimal("3000000"),
            registrado_por=self.admin,
        )
        content2 = _minimal_xlsx_rows(
            [{"numero": 1, "total": 4, "t1": 2, "t2": 1, "t3": 1, "valor": 12000000, "actividad": "Actualizada"}]
        )
        r2 = import_pic_excel(self.entity, 2026, content2, "pic2.xlsx", self.admin)
        self.assertTrue(r2["ok"])
        act.refresh_from_db()
        self.assertEqual(act.actividad, "Actualizada")
        self.assertEqual(act.ejecuciones.count(), 1)

    def test_tokenizar_encargado_multiple(self):
        tokens = tokenizar_encargado("ENFERMERA/ PSICOLGIA")
        self.assertIn("ENFERMERA", tokens)
        self.assertIn("PSICOLGIA", tokens)

    def test_auto_assign_from_cargo(self):
        contratista = User.objects.create_user(
            email="enf@test.com", password="x", entity=self.entity, is_active=True
        )
        contratista.groups.add(Group.objects.create(name="contratista"))
        UserEntityMembership.objects.create(
            user=contratista, entity=self.entity, role="contratista", is_active=True
        )
        PicCargo.objects.create(entity=self.entity, etiqueta="ENFERMERA", etiqueta_norm=normalizar_etiqueta("ENFERMERA"))
        cargo = PicCargo.objects.get(entity=self.entity, etiqueta_norm="ENFERMERA")
        cargo.usuarios.add(contratista)
        content = _minimal_xlsx_rows(
            [{"numero": 5, "total": 1, "t1": 1, "encargado": "ENFERMERA", "valor": 3000000}]
        )
        result = import_pic_excel(self.entity, 2026, content, "pic.xlsx", self.admin)
        self.assertTrue(result["ok"])
        act = PicActividad.objects.get(numero=5)
        self.assertIn(contratista.id, list(act.responsables.values_list("id", flat=True)))


class PicAccessTests(TestCase):
    def setUp(self):
        self.factory = APIRequestFactory()
        self.entity_a = Entity.objects.create(name="A", code="A", slug="ent-a", enable_pic=True)
        self.entity_b = Entity.objects.create(name="B", code="B", slug="ent-b", enable_pic=True)
        self.admin_group = Group.objects.create(name="admin")
        self.contra_group = Group.objects.create(name="contratista")

        self.admin_a = User.objects.create_user(email="a@test.com", password="x", entity=self.entity_a)
        self.admin_a.groups.add(self.admin_group)
        UserEntityMembership.objects.create(user=self.admin_a, entity=self.entity_a, role="admin", is_active=True)

        self.contra = User.objects.create_user(email="c@test.com", password="x", entity=self.entity_a)
        self.contra.groups.add(self.contra_group)
        UserEntityMembership.objects.create(
            user=self.contra, entity=self.entity_a, role="contratista", is_active=True, enabled_modules=["pic"]
        )

        plan = PicPlan.objects.create(entity=self.entity_a, anio=2026, uploaded_by=self.admin_a)
        self.act_asignada = PicActividad.objects.create(
            entity=self.entity_a,
            plan=plan,
            numero=1,
            actividad="Asignada",
            total_programado=4,
            prog_t1=1,
            prog_t2=1,
            prog_t3=1,
            prog_t4=1,
            valor_total=Decimal("12000000"),
            valor_unitario=Decimal("3000000"),
        )
        self.act_asignada.responsables.add(self.contra)

        self.act_otra = PicActividad.objects.create(
            entity=self.entity_a,
            plan=plan,
            numero=2,
            actividad="Otra",
            total_programado=1,
            prog_t1=1,
            valor_total=Decimal("1000000"),
            valor_unitario=Decimal("1000000"),
        )

    def test_contratista_only_sees_assigned(self):
        request = self.factory.get("/api/v1/pic/actividades/")
        force_authenticate(request, user=self.contra)
        view = PicActividadViewSet.as_view({"get": "list"})
        response = view(request)
        self.assertEqual(response.status_code, 200)
        ids = [item["id"] for item in response.data["results"]]
        self.assertIn(self.act_asignada.id, ids)
        self.assertNotIn(self.act_otra.id, ids)

    def test_ejecucion_exceeds_total_rejected(self):
        pdf = SimpleUploadedFile("ev.pdf", b"%PDF-1.4 test", content_type="application/pdf")
        request = self.factory.post(
            f"/api/v1/pic/actividades/{self.act_asignada.id}/ejecuciones/",
            {
                "fecha_ejecucion": "2026-03-01",
                "cantidad_ejecutada": "5",
                "descripcion": "Exceso",
                "archivos": pdf,
            },
            format="multipart",
        )
        force_authenticate(request, user=self.contra)
        view = PicActividadViewSet.as_view({"post": "ejecuciones"})
        response = view(request, pk=self.act_asignada.id)
        self.assertEqual(response.status_code, 400)
