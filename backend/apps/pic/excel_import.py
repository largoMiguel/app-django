"""Importación Excel de seguimiento PIC."""
from __future__ import annotations

import re
import unicodedata
from decimal import Decimal, InvalidOperation
from io import BytesIO
from typing import Any

from django.db import transaction
from openpyxl import load_workbook

from .calculos import calcular_valor_unitario
from .models import PicActividad, PicCargo, PicPlan


def normalizar_etiqueta(texto: str) -> str:
    raw = (texto or "").strip().upper()
    raw = unicodedata.normalize("NFKD", raw)
    raw = "".join(c for c in raw if not unicodedata.combining(c))
    raw = re.sub(r"\s+", " ", raw)
    return raw.strip()


def tokenizar_encargado(texto: str) -> list[str]:
    if not texto or not str(texto).strip():
        return []
    parts = re.split(r"[/,]", str(texto))
    tokens = [normalizar_etiqueta(p) for p in parts if normalizar_etiqueta(p)]
    return list(dict.fromkeys(tokens))


def _cell_value(cell) -> Any:
    if cell is None:
        return None
    return cell.value


def _to_int(value, default: int = 0) -> int:
    if value is None or value == "":
        return default
    try:
        return int(round(float(value)))
    except (TypeError, ValueError):
        return default


def _to_decimal(value) -> Decimal:
    if value is None or value == "":
        return Decimal("0")
    try:
        if isinstance(value, str):
            value = value.replace(",", "").strip()
        return Decimal(str(value))
    except (InvalidOperation, ValueError):
        return Decimal("0")


def _to_text(value) -> str:
    if value is None:
        return ""
    return str(value).strip()


def parse_pic_excel(file_content: bytes) -> tuple[list[dict], list[str]]:
    """Parsea el Excel PIC. Retorna (filas, errores_globales)."""
    wb = load_workbook(filename=BytesIO(file_content), data_only=True, read_only=True)
    ws = wb.active
    if ws is None:
        return [], ["El archivo no contiene hojas."]

    rows: list[dict] = []
    numeros_vistos: dict[int, int] = {}
    errores: list[str] = []

    for row_idx, row in enumerate(ws.iter_rows(min_row=3), start=3):
        cells = {chr(64 + i): _cell_value(c) for i, c in enumerate(row[:14], start=1)}
        numero_raw = cells.get("C")
        if numero_raw is None or str(numero_raw).strip() == "":
            continue

        numero = _to_int(numero_raw, default=-1)
        if numero <= 0:
            errores.append(f"Fila {row_idx}: número de actividad inválido ({numero_raw}).")
            continue

        if numero in numeros_vistos:
            errores.append(
                f"Fila {row_idx}: número {numero} duplicado (también en fila {numeros_vistos[numero]})."
            )
        numeros_vistos[numero] = row_idx

        total_programado = _to_int(cells.get("H"))
        prog_t1 = _to_int(cells.get("I"))
        prog_t2 = _to_int(cells.get("J"))
        prog_t3 = _to_int(cells.get("K"))
        prog_t4 = _to_int(cells.get("L"))
        valor_total = _to_decimal(cells.get("M"))

        if total_programado <= 0:
            errores.append(f"Fila {row_idx} (actividad {numero}): total programado debe ser mayor a 0.")
            continue

        rows.append(
            {
                "fila_excel": row_idx,
                "numero": numero,
                "eje_estrategico": _to_text(cells.get("A")),
                "linea_operativa": _to_text(cells.get("B")),
                "encargado_texto": _to_text(cells.get("D")),
                "actividad": _to_text(cells.get("E")),
                "soportes": _to_text(cells.get("F")),
                "unidad_medida": _to_text(cells.get("G")),
                "total_programado": total_programado,
                "prog_t1": prog_t1,
                "prog_t2": prog_t2,
                "prog_t3": prog_t3,
                "prog_t4": prog_t4,
                "valor_total": valor_total,
                "valor_unitario": calcular_valor_unitario(valor_total, total_programado),
            }
        )

    wb.close()
    return rows, errores


def _usuarios_para_encargado(entity, encargado_texto: str) -> tuple[list, list[str]]:
    """Resuelve usuarios desde PicCargo. Retorna (user_ids, tokens_sin_mapeo)."""
    tokens = tokenizar_encargado(encargado_texto)
    if not tokens:
        return [], []

    usuarios_ids: set[int] = set()
    sin_mapeo: list[str] = []
    cargos = {
        c.etiqueta_norm: c
        for c in PicCargo.objects.filter(entity=entity).prefetch_related("usuarios")
    }
    for token in tokens:
        cargo = cargos.get(token)
        if cargo is None:
            sin_mapeo.append(token)
            continue
        usuarios_ids.update(cargo.usuarios.values_list("id", flat=True))
    return list(usuarios_ids), sin_mapeo


def import_pic_excel(
    entity,
    anio: int,
    file_content: bytes,
    filename: str,
    user,
    *,
    reasignar_encargados: bool = False,
) -> dict:
    """Importa o actualiza el plan PIC desde Excel."""
    rows, errores = parse_pic_excel(file_content)
    if errores:
        return {"ok": False, "errores": errores}

    if not rows:
        return {"ok": False, "errores": ["No se encontraron actividades en el Excel."]}

    advertencias: list[str] = []
    sin_mapeo_global: set[str] = set()
    creadas = 0
    actualizadas = 0
    eliminadas = 0
    conservadas = 0

    with transaction.atomic():
        plan, _ = PicPlan.objects.get_or_create(
            entity=entity,
            anio=anio,
            defaults={"uploaded_by": user},
        )
        plan.archivo_nombre = filename
        plan.uploaded_by = user
        plan.save(update_fields=["archivo_nombre", "uploaded_by", "updated_at"])

        numeros_excel = {r["numero"] for r in rows}
        existing = {a.numero: a for a in PicActividad.objects.filter(plan=plan)}

        for row in rows:
            suma_trim = row["prog_t1"] + row["prog_t2"] + row["prog_t3"] + row["prog_t4"]
            if suma_trim != row["total_programado"]:
                advertencias.append(
                    f"Actividad {row['numero']}: suma trimestres ({suma_trim}) "
                    f"≠ total programado ({row['total_programado']})."
                )

            act = existing.get(row["numero"])
            is_new = act is None
            if is_new:
                act = PicActividad(entity=entity, plan=plan, numero=row["numero"])
                creadas += 1
            else:
                actualizadas += 1

            for field in (
                "fila_excel",
                "eje_estrategico",
                "linea_operativa",
                "encargado_texto",
                "actividad",
                "soportes",
                "unidad_medida",
                "total_programado",
                "prog_t1",
                "prog_t2",
                "prog_t3",
                "prog_t4",
                "valor_total",
                "valor_unitario",
            ):
                setattr(act, field, row[field])
            act.save()

            user_ids, sin_mapeo = _usuarios_para_encargado(entity, row["encargado_texto"])
            sin_mapeo_global.update(sin_mapeo)

            tiene_responsables = act.responsables.exists()
            if is_new or reasignar_encargados or not tiene_responsables:
                if user_ids:
                    act.responsables.set(user_ids)
                    first_cargo = PicCargo.objects.filter(
                        entity=entity,
                        etiqueta_norm__in=tokenizar_encargado(row["encargado_texto"]),
                        secretaria__isnull=False,
                    ).first()
                    if first_cargo and first_cargo.secretaria_id:
                        act.responsable_secretaria_id = first_cargo.secretaria_id
                        act.save(update_fields=["responsable_secretaria_id", "updated_at"])

        for numero, act in existing.items():
            if numero in numeros_excel:
                continue
            if act.ejecuciones.exists():
                conservadas += 1
                advertencias.append(
                    f"Actividad {numero} ya no está en el Excel pero tiene ejecuciones; se conservó."
                )
                continue
            act.delete()
            eliminadas += 1

    valor_total_pic = sum((r["valor_total"] for r in rows), Decimal("0"))

    return {
        "ok": True,
        "anio": anio,
        "plan_id": plan.id,
        "total_actividades": len(rows),
        "creadas": creadas,
        "actualizadas": actualizadas,
        "eliminadas": eliminadas,
        "conservadas_sin_excel": conservadas,
        "sin_mapeo_encargado": sorted(sin_mapeo_global),
        "advertencias": advertencias,
        "valor_total_pic": str(valor_total_pic),
        "fecha_carga": plan.updated_at.isoformat() if plan.updated_at else None,
    }


def aplicar_cargos_a_actividades(entity, plan: PicPlan) -> dict:
    """Reasigna responsables en bloque según mapeo PicCargo."""
    actualizadas = 0
    sin_mapeo: set[str] = set()
    with transaction.atomic():
        for act in PicActividad.objects.filter(plan=plan, entity=entity):
            user_ids, tokens = _usuarios_para_encargado(entity, act.encargado_texto)
            sin_mapeo.update(tokens)
            if user_ids:
                act.responsables.set(user_ids)
                actualizadas += 1
    return {"actualizadas": actualizadas, "sin_mapeo_encargado": sorted(sin_mapeo)}
