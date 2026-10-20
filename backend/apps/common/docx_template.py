"""Membrete institucional en Word (.docx) → PDF vía LibreOffice."""
from __future__ import annotations

import logging
import os
import re
import shutil
import subprocess
import tempfile
import zipfile
from copy import deepcopy
from dataclasses import dataclass
from io import BytesIO
from xml.etree import ElementTree as ET

import fitz
from django.conf import settings
from docx import Document
from docx.enum.text import WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

from apps.common.b2_client import get_b2_client
from apps.common.pdf_template import apply_template_overlay, detect_template_margins

logger = logging.getLogger(__name__)

_PLAIN_PAGINATION_RE = re.compile(
    r"P[aá]gina\s+\d+\s+de\s+\d+",
    re.IGNORECASE,
)
_MARKER_PAGINA = "{{PAGINA}}"
_MARKER_TOTAL = "{{TOTAL}}"

W_NS = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
_W_TAG = f"{{{W_NS}}}"

# Fuente solicitada en Word → familias instaladas que la sustituyen con métricas compatibles.
_FONT_EQUIVALENTS: dict[str, frozenset[str]] = {
    "arial": frozenset({"arial", "liberation sans", "helvetica", "nimbus sans l"}),
    "helvetica": frozenset({"helvetica", "liberation sans", "arial", "nimbus sans l"}),
    "times new roman": frozenset({"times new roman", "liberation serif", "nimbus roman"}),
    "courier new": frozenset({"courier new", "liberation mono", "nimbus mono l"}),
    "arial narrow": frozenset({"arial narrow", "liberation sans narrow"}),
    "calibri": frozenset({"calibri", "carlito"}),
    "cambria": frozenset({"cambria", "caladea"}),
    "century gothic": frozenset({"century gothic", "urw gothic", "urw gothic l"}),
    "verdana": frozenset({"verdana", "dejavu sans"}),
    "tahoma": frozenset({"tahoma", "dejavu sans condensed"}),
    "palatino linotype": frozenset({"palatino linotype", "p052", "urw palladio l"}),
    "bookman old style": frozenset({"bookman old style", "urw bookman l"}),
    "century schoolbook": frozenset({"century schoolbook", "century schoolbook l", "nimbus roman"}),
}


def _normalize_font_key(name: str) -> str:
    return " ".join(name.strip().lower().split())


def collect_requested_fonts(docx_bytes: bytes) -> set[str]:
    """Familias tipográficas referenciadas en el .docx (runs y estilos)."""
    families: set[str] = set()
    with zipfile.ZipFile(BytesIO(docx_bytes)) as zf:
        for name in zf.namelist():
            if not name.startswith("word/") or not name.endswith(".xml"):
                continue
            if name.endswith("fontTable.xml"):
                continue
            try:
                root = ET.fromstring(zf.read(name))
            except ET.ParseError:
                continue
            for node in root.iter(f"{_W_TAG}rFonts"):
                for attr in ("ascii", "hAnsi", "cs", "eastAsia"):
                    val = node.get(f"{_W_TAG}{attr}")
                    if val and val.strip():
                        families.add(val.strip())
    return families


_fc_available_cache: set[str] | None = None


def invalidate_fc_available_cache() -> None:
    global _fc_available_cache
    _fc_available_cache = None


def _fc_available_families() -> set[str]:
    global _fc_available_cache
    if _fc_available_cache is not None:
        return _fc_available_cache
    try:
        proc = subprocess.run(
            ["fc-list", "--format=%{family}\n"],
            capture_output=True,
            text=True,
            timeout=30,
            check=False,
        )
    except (FileNotFoundError, subprocess.TimeoutExpired):
        _fc_available_cache = set()
        return _fc_available_cache
    if proc.returncode != 0:
        _fc_available_cache = set()
        return _fc_available_cache
    out: set[str] = set()
    for line in proc.stdout.splitlines():
        for part in line.split(","):
            key = _normalize_font_key(part)
            if key:
                out.add(key)
    _fc_available_cache = out
    return out


def _font_is_available(requested: str, available: set[str]) -> bool:
    key = _normalize_font_key(requested)
    if key in available:
        return True
    equivalents = _FONT_EQUIVALENTS.get(key, frozenset({key}))
    return bool(equivalents & available)


def missing_fonts(requested: set[str] | list[str]) -> list[str]:
    """Fuentes del documento sin sustituto métrico conocido en el servidor."""
    available = _fc_available_families()
    missing: list[str] = []
    for font in sorted(requested, key=str.lower):
        if not _font_is_available(font, available):
            missing.append(font)
    return missing


def analyze_docx_fonts(
    docx_bytes: bytes,
    *,
    auto_install: bool | None = None,
) -> tuple[list[str], list[str], list[str]]:
    """Devuelve (fuentes solicitadas, aún faltantes, instaladas en esta pasada)."""
    if auto_install is None:
        auto_install = bool(getattr(settings, "DOCX_AUTO_INSTALL_FONTS", True))
    requested_set = collect_requested_fonts(docx_bytes)
    installed: list[str] = []
    if auto_install and requested_set:
        from apps.common.font_install import ensure_fonts_for_docx

        result = ensure_fonts_for_docx(docx_bytes, requested_set)
        installed = result.installed
    requested = sorted(requested_set, key=str.lower)
    missing = missing_fonts(requested)
    return requested, missing, installed


def load_entity_docx_template(entity) -> bytes | None:
    """Descarga el membrete .docx desde B2 (bucket PQRS)."""
    url = getattr(entity, "docx_template_url", None)
    if not url:
        return None
    try:
        key = url.lstrip("/")
        client = get_b2_client()
        resp = client.get_object(Bucket=settings.B2_BUCKET_PQRS, Key=key)
        return resp["Body"].read()
    except Exception:
        logger.exception(
            "Error descargando template DOCX para entidad %s",
            getattr(entity, "id", None),
        )
        return None


def _paragraph_has_page_fields(paragraph) -> bool:
    xml = paragraph._element.xml
    if "fldSimple" not in xml and "fldChar" not in xml:
        return False
    return "PAGE" in xml.upper() or "NUMPAGES" in xml.upper()


def _first_run_rpr(paragraph):
    for run in paragraph.runs:
        r_pr = run._element.find(qn("w:rPr"))
        if r_pr is not None:
            return deepcopy(r_pr)
    return None


def _make_run(text: str, r_pr=None) -> OxmlElement:
    r = OxmlElement("w:r")
    if r_pr is not None:
        r.append(deepcopy(r_pr))
    t = OxmlElement("w:t")
    t.set(qn("xml:space"), "preserve")
    t.text = text
    r.append(t)
    return r


def _make_fld_simple(instr: str, r_pr=None) -> OxmlElement:
    fld = OxmlElement("w:fldSimple")
    fld.set(qn("w:instr"), f" {instr.strip()} ")
    r = OxmlElement("w:r")
    if r_pr is not None:
        r.append(deepcopy(r_pr))
    t = OxmlElement("w:t")
    t.text = "1"
    r.append(t)
    fld.append(r)
    return fld


def _clear_paragraph_runs(paragraph) -> None:
    p = paragraph._element
    for child in list(p):
        if child.tag in (qn("w:r"), qn("w:hyperlink"), qn("w:fldSimple"), qn("w:smartTag")):
            p.remove(child)


def _rebuild_pagination_paragraph(paragraph, *, use_markers: bool) -> bool:
    if _paragraph_has_page_fields(paragraph):
        return False

    full_text = paragraph.text
    if use_markers:
        if _MARKER_PAGINA not in full_text and _MARKER_TOTAL not in full_text:
            return False
    elif not _PLAIN_PAGINATION_RE.search(full_text):
        return False

    r_pr = _first_run_rpr(paragraph)
    _clear_paragraph_runs(paragraph)
    p = paragraph._element

    if use_markers:
        parts = re.split(
            rf"({re.escape(_MARKER_PAGINA)}|{re.escape(_MARKER_TOTAL)})",
            full_text,
        )
        for part in parts:
            if not part:
                continue
            if part == _MARKER_PAGINA:
                p.append(_make_fld_simple("PAGE", r_pr))
            elif part == _MARKER_TOTAL:
                p.append(_make_fld_simple("NUMPAGES", r_pr))
            else:
                p.append(_make_run(part, r_pr))
        return True

    p.append(_make_run("Página ", r_pr))
    p.append(_make_fld_simple("PAGE", r_pr))
    p.append(_make_run(" de ", r_pr))
    p.append(_make_fld_simple("NUMPAGES", r_pr))
    return True


def _iter_paragraphs_in_container(container):
    for paragraph in container.paragraphs:
        yield paragraph
    for table in container.tables:
        for row in table.rows:
            for cell in row.cells:
                yield from _iter_paragraphs_in_container(cell)


def _iter_header_footer_parts(document):
    for section in document.sections:
        for part in (
            section.header,
            section.footer,
            section.first_page_header,
            section.first_page_footer,
            section.even_page_header,
            section.even_page_footer,
        ):
            if part is None:
                continue
            try:
                if part.is_linked_to_previous:
                    continue
            except Exception:
                pass
            yield part


def normalize_pagination_fields(document: Document) -> None:
    """Convierte texto plano o marcadores en campos PAGE / NUMPAGES."""
    for part in _iter_header_footer_parts(document):
        for paragraph in _iter_paragraphs_in_container(part):
            text = paragraph.text
            if _MARKER_PAGINA in text or _MARKER_TOTAL in text:
                _rebuild_pagination_paragraph(paragraph, use_markers=True)
            elif _PLAIN_PAGINATION_RE.search(text):
                _rebuild_pagination_paragraph(paragraph, use_markers=False)


def append_page_breaks(document: Document, extra_pages: int) -> None:
    """Añade saltos de página al final del cuerpo."""
    if extra_pages <= 0:
        return
    for _ in range(extra_pages):
        paragraph = document.add_paragraph()
        run = paragraph.add_run()
        run.add_break(WD_BREAK.PAGE)


def _find_soffice() -> str | None:
    for name in ("soffice", "libreoffice"):
        path = shutil.which(name)
        if path:
            return path
    return None


def _convert_with_libreoffice(docx_path: str, outdir: str) -> str:
    soffice = _find_soffice()
    if not soffice:
        raise RuntimeError("LibreOffice (soffice) no está instalado")

    profile_dir = tempfile.mkdtemp(prefix="lo_profile_")
    profile_url = "file://" + profile_dir.replace("\\", "/")
    timeout = int(getattr(settings, "LIBREOFFICE_TIMEOUT", 120))
    cmd = [
        soffice,
        "--headless",
        f"-env:UserInstallation={profile_url}",
        "--convert-to",
        "pdf",
        "--outdir",
        outdir,
        docx_path,
    ]
    try:
        subprocess.run(
            cmd,
            check=True,
            timeout=timeout,
            capture_output=True,
            text=True,
        )
    finally:
        shutil.rmtree(profile_dir, ignore_errors=True)

    base = os.path.splitext(os.path.basename(docx_path))[0]
    pdf_path = os.path.join(outdir, f"{base}.pdf")
    if not os.path.isfile(pdf_path):
        raise RuntimeError(f"LibreOffice no generó PDF en {pdf_path}")
    return pdf_path


def render_template_pdf(docx_bytes: bytes, total_pages: int) -> bytes:
    """Genera PDF del membrete con total_pages hojas y paginación resuelta."""
    if getattr(settings, "DOCX_AUTO_INSTALL_FONTS", True):
        from apps.common.font_install import ensure_fonts_for_docx

        ensure_fonts_for_docx(docx_bytes)
    document = Document(BytesIO(docx_bytes))
    normalize_pagination_fields(document)
    append_page_breaks(document, max(0, total_pages - 1))

    with tempfile.TemporaryDirectory(prefix="docx_tpl_") as tmp:
        docx_path = os.path.join(tmp, "template.docx")
        document.save(docx_path)
        pdf_path = _convert_with_libreoffice(docx_path, tmp)
        with open(pdf_path, "rb") as fh:
            return fh.read()


@dataclass
class PreparedEntityTemplate:
    docx_bytes: bytes
    top_margin_in: float
    bottom_margin_in: float

    def render_overlay(self, content_pdf_bytes: bytes) -> BytesIO:
        doc = fitz.open(stream=content_pdf_bytes, filetype="pdf")
        total = len(doc)
        doc.close()
        template_pdf = render_template_pdf(self.docx_bytes, total)
        return apply_template_overlay(content_pdf_bytes, template_pdf)


def prepare_entity_template(entity) -> PreparedEntityTemplate | None:
    """Carga membrete .docx y detecta márgenes desde una conversión de prueba."""
    docx_bytes = load_entity_docx_template(entity)
    if not docx_bytes:
        return None

    top_in, bottom_in = 0.8, 0.8
    try:
        sample_pdf = render_template_pdf(docx_bytes, 1)
        top_in, bottom_in = detect_template_margins(sample_pdf)
    except Exception:
        logger.exception(
            "No se pudieron detectar márgenes del membrete DOCX (entidad %s)",
            getattr(entity, "id", None),
        )

    return PreparedEntityTemplate(
        docx_bytes=docx_bytes,
        top_margin_in=top_in,
        bottom_margin_in=bottom_in,
    )
