"""Pruebas de membrete Word (.docx) para informes."""
from __future__ import annotations

import shutil
from io import BytesIO

import fitz
import pytest
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

from apps.common.docx_template import (
    _find_soffice,
    append_page_breaks,
    collect_requested_fonts,
    missing_fonts,
    normalize_pagination_fields,
)
from apps.common.pdf_template import apply_template_overlay


def _docx_with_font(font_name: str) -> bytes:
    doc = Document()
    paragraph = doc.add_paragraph()
    run = paragraph.add_run("Texto de prueba")
    rfonts = OxmlElement("w:rFonts")
    rfonts.set(qn("w:ascii"), font_name)
    rfonts.set(qn("w:hAnsi"), font_name)
    run._element.get_or_add_rPr().append(rfonts)
    buf = BytesIO()
    doc.save(buf)
    return buf.getvalue()


def _docx_with_header_text(text: str) -> bytes:
    doc = Document()
    doc.sections[0].header.paragraphs[0].text = text
    buf = BytesIO()
    doc.save(buf)
    return buf.getvalue()


def test_normalize_plain_pagination_to_fields():
    doc = Document(BytesIO(_docx_with_header_text("Página 1 de 12")))
    normalize_pagination_fields(doc)
    xml = doc.sections[0].header.paragraphs[0]._element.xml
    assert "fldSimple" in xml
    assert "PAGE" in xml.upper()
    assert "NUMPAGES" in xml.upper()
    assert "Página 1 de 12" not in doc.sections[0].header.paragraphs[0].text


def test_normalize_markers_to_fields():
    doc = Document(
        BytesIO(_docx_with_header_text("Página {{PAGINA}} de {{TOTAL}}"))
    )
    normalize_pagination_fields(doc)
    xml = doc.sections[0].header.paragraphs[0]._element.xml
    assert "PAGE" in xml.upper()
    assert "NUMPAGES" in xml.upper()
    assert "{{PAGINA}}" not in xml


def test_append_page_breaks_count():
    doc = Document()
    append_page_breaks(doc, 3)
    xml = doc._element.body.xml
    assert xml.count('w:type="page"') == 3


def test_apply_template_overlay_multipage():
    tpl = fitz.open()
    for label in ("A", "B", "C"):
        page = tpl.new_page(width=595, height=842)
        page.insert_text((72, 72), f"Membrete {label}", fontsize=14)
    tpl_bytes = tpl.tobytes()
    tpl.close()

    content = fitz.open()
    for _ in range(3):
        content.new_page(width=595, height=842)
    content_bytes = content.tobytes()
    content.close()

    out = apply_template_overlay(content_bytes, tpl_bytes)
    doc = fitz.open(stream=out.getvalue(), filetype="pdf")
    assert len(doc) == 3
    assert "Membrete A" in doc[0].get_text()
    assert "Membrete B" in doc[1].get_text()
    assert "Membrete C" in doc[2].get_text()
    doc.close()


def test_collect_requested_fonts_detects_arial():
    docx_bytes = _docx_with_font("Arial")
    assert "Arial" in collect_requested_fonts(docx_bytes)


def test_missing_fonts_flags_unknown(monkeypatch):
    monkeypatch.setattr(
        "apps.common.docx_template._fc_available_families",
        lambda: {"liberation sans", "dejavu sans"},
    )
    missing = missing_fonts({"FuenteInexistenteXYZ"})
    assert missing == ["FuenteInexistenteXYZ"]


def test_missing_fonts_arial_has_substitute(monkeypatch):
    monkeypatch.setattr(
        "apps.common.docx_template._fc_available_families",
        lambda: {"liberation sans"},
    )
    assert missing_fonts({"Arial"}) == []


@pytest.mark.skipif(_find_soffice() is None, reason="LibreOffice no instalado")
def test_render_template_pdf_integration():
    from apps.common.docx_template import render_template_pdf

    docx_bytes = _docx_with_header_text("Página {{PAGINA}} de {{TOTAL}}")
    pdf_bytes = render_template_pdf(docx_bytes, 2)
    doc = fitz.open(stream=pdf_bytes, filetype="pdf")
    assert len(doc) >= 2
    doc.close()
